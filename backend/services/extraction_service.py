import asyncio
import logging
from io import BytesIO

logger = logging.getLogger(__name__)


def _extract_pdf(content: bytes) -> str:
    import fitz  # PyMuPDF — pip install PyMuPDF; import name is fitz
    doc = fitz.open(stream=content, filetype="pdf")
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    return text


def _extract_docx(content: bytes) -> str:
    from docx import Document  # pip install python-docx; import name is docx
    doc = Document(BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cell_text = " | ".join(c.text for c in row.cells if c.text.strip())
            if cell_text:
                parts.append(cell_text)
    return "\n".join(parts)


async def extract_text(content: bytes, content_type: str) -> str:
    loop = asyncio.get_running_loop()
    if "pdf" in content_type:
        return await loop.run_in_executor(None, _extract_pdf, content)
    else:
        return await loop.run_in_executor(None, _extract_docx, content)
